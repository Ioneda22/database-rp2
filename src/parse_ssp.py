"""
Lê os arquivos da SSP-SP e agrega as ocorrências por município, ano e mês.

Saídas em data/processed/:
  ssp_painel.csv       codigo_ibge, ano, mes, natureza, ocorrencias
  ssp_textura.csv      codigo_ibge, ano, contagens de local e período
  ssp_complementar.csv codigo_ibge, ano, veículos e celulares
  ssp_cobertura.csv    ano, quantos meses tem

"""
from __future__ import annotations

import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from config import (
    SSP_DIR, SSP_CRIMINAIS_GLOB, SSP_VEICULOS_GLOB, SSP_CELULARES_GLOB,
    SSP_PAINEL_CSV, SSP_TEXTURA_CSV, SSP_COMPLEMENTAR_CSV, SSP_COBERTURA_CSV,
)


def normaliza(txt) -> str:
    """Deixa o texto em maiúsculas, sem acento e sem espaços repetidos."""
    txt = str(txt).strip().upper()
    txt = unicodedata.normalize("NFKD", txt).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", txt)

APELIDOS_CRIMINAIS = {
    "cod_ibge":  ["CD_IBGE", "COD IBGE"],
    "ano":       ["ANO_ESTATISTICA"],
    "mes":       ["MES_ESTATISTICA"],
    "natureza":  ["NATUREZA_APURADA"],
    "local":     ["DESCR_SUBTIPOLOCAL"],
    "periodo":   ["DESC_PERIODO", "DESCR_PERIODO"],
}

APELIDOS_VEICULOS = {
    "cod_ibge": ["CD_IBGE", "COD IBGE"],
    "ano":      ["ANO_REGISTRO_BO", "ANO"],
    "desfecho": ["DESCR_OCORRENCIA_VEICULO"],
    "tipo":     ["DESCR_TIPO_VEICULO"],
}

APELIDOS_CELULARES = {
    "cod_ibge": ["CD_IBGE", "COD IBGE"],
    "ano":      ["ANO_REGISTRO_BO", "ANO"],
    "rubrica":  ["RUBRICA"],
}


def _indices(cabecalho: list,
             apelidos: dict[str, list[str]]) -> dict[str, int | None]:
    """
    Descobre em qual posição da linha está cada campo, aceitando qualquer
    um dos nomes da lista. Campos que não existirem viram None.
    """
    posicoes = {}
    for i, c in enumerate(cabecalho):
        if c is not None:
            posicoes.setdefault(normaliza(c), i)

    return {campo: next((posicoes[normaliza(n)] for n in nomes
                         if normaliza(n) in posicoes), None)
            for campo, nomes in apelidos.items()}


def _abas_de_dados(wb, apelidos: dict[str, list[str]],
                   marcadores: list[str]) -> list[str]:
    """
    Descobre quais abas têm dados olhando o cabeçalho, e não o nome da aba
    (o nome muda de ano para ano). Uma aba serve se tiver todos os campos
    de `marcadores`.
    """
    abas = []
    for nome in wb.sheetnames:
        ws = wb[nome]
        try:
            cabecalho = next(ws.iter_rows(max_row=1, values_only=True))
        except StopIteration:
            continue
        presentes = {normaliza(c) for c in cabecalho if c is not None}
        if all(any(normaliza(n) in presentes for n in apelidos[campo])
               for campo in marcadores):
            abas.append(nome)
    return abas


def _cod_ibge(valor) -> str | None:
    """Devolve o código IBGE como texto de 7 dígitos, ou None se inválido."""
    if valor is None:
        return None
    txt = str(valor).strip()
    if txt in ("", "NULL"):
        return None
    if txt.endswith(".0"):
        txt = txt[:-2]
    return txt if txt.isdigit() and len(txt) == 7 else None


# ---------------------------------------------------------------------------
# 1. Arquivo principal: SPDadosCriminais
# ---------------------------------------------------------------------------

PERIODOS_NOTURNOS = {"A NOITE", "DE MADRUGADA"}
PERIODOS_DIURNOS = {"PELA MANHA", "A TARDE"}


def agregar_criminais(
    diretorio: Path = SSP_DIR,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """
    Lê todos os SPDadosCriminais_*.xlsx e devolve (painel, textura,
    cobertura).

    O ano é o ANO_ESTATISTICA (quando a ocorrência entrou na estatística),
    e não o ANO_BO, porque é assim que a SSP conta.
    """
    arquivos = sorted(diretorio.glob(SSP_CRIMINAIS_GLOB))

    contagem = Counter()                    # (cod, ano, mes, natureza) -> n
    textura = defaultdict(Counter)          # (cod, ano) -> contadores
    meses = defaultdict(set)                # ano -> meses que apareceram

    for arq in arquivos:
        wb = load_workbook(arq, read_only=True, data_only=True)
        abas = _abas_de_dados(wb, APELIDOS_CRIMINAIS, ["cod_ibge", "natureza"])
        if not abas:
            wb.close()
            continue

        for aba in abas:
            ws = wb[aba]
            it = ws.iter_rows(values_only=True)
            idx = _indices(list(next(it)), APELIDOS_CRIMINAIS)
            i_cod, i_ano, i_mes = idx["cod_ibge"], idx["ano"], idx["mes"]
            i_nat, i_loc, i_per = idx["natureza"], idx["local"], idx["periodo"]

            for row in it:
                cod = _cod_ibge(row[i_cod])
                if cod is None or row[i_ano] is None:
                    continue
                ano = int(row[i_ano])
                chave = (cod, ano)
                natureza = normaliza(row[i_nat])
                # Se o mês estiver vazio, fica 0. A linha continua contando
                # no total do ano; só a série mensal filtra mes > 0.
                mes = int(row[i_mes]) if row[i_mes] is not None else 0
                contagem[(cod, ano, mes, natureza)] += 1

                t = textura[chave]
                t["n_ocorrencias"] += 1
                if i_loc is not None and normaliza(row[i_loc]) == "VIA PUBLICA":
                    t["n_via_publica"] += 1
                if i_per is not None:
                    periodo = normaliza(row[i_per])
                    if periodo in PERIODOS_NOTURNOS:
                        t["n_periodo_conhecido"] += 1
                        t["n_noturno"] += 1
                    elif periodo in PERIODOS_DIURNOS:
                        t["n_periodo_conhecido"] += 1

                if mes:
                    meses[ano].add(mes)
        wb.close()

    painel = pd.DataFrame(
        [(c, a, m, n, q) for (c, a, m, n), q in sorted(contagem.items())],
        columns=["codigo_ibge", "ano", "mes", "natureza", "ocorrencias"],
    )
    tex = pd.DataFrame(
        [
            {"codigo_ibge": c, "ano": a,
             "n_ocorrencias": v["n_ocorrencias"],
             "n_via_publica": v["n_via_publica"],
             "n_periodo_conhecido": v["n_periodo_conhecido"],
             "n_noturno": v["n_noturno"]}
            for (c, a), v in sorted(textura.items())
        ]
    )
    cob = pd.DataFrame(
        [{"ano": a, "n_meses": len(m),
          "meses": ",".join(str(x) for x in sorted(m))}
         for a, m in sorted(meses.items())]
    )
    return painel, tex, cob


# ---------------------------------------------------------------------------
# 2. Arquivos de veículos e celulares
# ---------------------------------------------------------------------------

MOTOS = {"MOTOCICLO", "MOTONETA", "CICLOMOTO"}


def agregar_complementares(diretorio: Path = SSP_DIR) -> pd.DataFrame:
    """Agrega VeiculosSubtraidos_* e CelularesSubtraidos_* por município e ano."""
    acc = defaultdict(Counter)   # (cod, ano) -> contadores

    # --- veículos ---
    for arq in sorted(diretorio.glob(SSP_VEICULOS_GLOB)):
        wb = load_workbook(arq, read_only=True, data_only=True)
        for aba in _abas_de_dados(wb, APELIDOS_VEICULOS,
                                  ["cod_ibge", "desfecho"]):
            ws = wb[aba]
            it = ws.iter_rows(values_only=True)
            idx = _indices(list(next(it)), APELIDOS_VEICULOS)
            for row in it:
                cod = _cod_ibge(row[idx["cod_ibge"]])
                ano = row[idx["ano"]]
                desfecho = row[idx["desfecho"]]
                # Alguns arquivos têm linhas vazias no final.
                if cod is None or ano is None or desfecho is None:
                    continue
                chave = (cod, int(ano))
                d = normaliza(desfecho)
                if d == "FURTADO":
                    acc[chave]["veic_furtado"] += 1
                elif d == "ROUBADO":
                    acc[chave]["veic_roubado"] += 1
                elif d.startswith("LOCALIZADO"):
                    # "Localizado / Entregue" é veículo recuperado, não crime.
                    acc[chave]["veic_localizado"] += 1
                if d in ("FURTADO", "ROUBADO"):
                    acc[chave]["veic_subtraido"] += 1
                    if normaliza(row[idx["tipo"]]) in MOTOS:
                        acc[chave]["veic_moto"] += 1
        wb.close()

    # --- celulares ---
    for arq in sorted(diretorio.glob(SSP_CELULARES_GLOB)):
        wb = load_workbook(arq, read_only=True, data_only=True)
        for aba in _abas_de_dados(wb, APELIDOS_CELULARES,
                                  ["cod_ibge", "rubrica"]):
            ws = wb[aba]
            it = ws.iter_rows(values_only=True)
            idx = _indices(list(next(it)), APELIDOS_CELULARES)
            for row in it:
                cod = _cod_ibge(row[idx["cod_ibge"]])
                ano = row[idx["ano"]]
                if cod is None or ano is None:
                    continue
                rubrica = normaliza(row[idx["rubrica"]])
                # "Perda/Extravio" não é crime.
                if rubrica.startswith("FURTO"):
                    acc[(cod, int(ano))]["cel_furto"] += 1
                elif rubrica.startswith("ROUBO"):
                    acc[(cod, int(ano))]["cel_roubo"] += 1
        wb.close()

    campos = ["veic_furtado", "veic_roubado", "veic_localizado",
              "veic_subtraido", "veic_moto", "cel_furto", "cel_roubo"]
    return pd.DataFrame(
        [dict({"codigo_ibge": c, "ano": a}, **{k: v[k] for k in campos})
         for (c, a), v in sorted(acc.items())]
    )


def main() -> None:
    painel, textura, cobertura = agregar_criminais()
    painel.to_csv(SSP_PAINEL_CSV, index=False, encoding="utf-8")
    textura.to_csv(SSP_TEXTURA_CSV, index=False, encoding="utf-8")
    cobertura.to_csv(SSP_COBERTURA_CSV, index=False, encoding="utf-8")

    comp = agregar_complementares()
    comp.to_csv(SSP_COMPLEMENTAR_CSV, index=False, encoding="utf-8")


if __name__ == "__main__":
    main()
